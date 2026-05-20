import openpyxl
import os

file_path = r'data\actual\excel_sources\S1CRUF_input.xlsx'
try:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    print(f"File loaded. Sheets: {wb.sheetnames}")
    for sn in wb.sheetnames:
        ws = wb[sn]
        print(f"Scanning sheet: {sn}")
        for row in ws.iter_rows(max_row=100, max_col=20): # Limit scan to first 100x20
             for cell in row:
                 if cell.value:
                     s_val = str(cell.value)
                     if any(kw in s_val for kw in ['DIMINPUTCODE', 'OUTPUTMCCOMP', 'INPUTMCCOMP']):
                         print(f"Found match: {cell.coordinate} in {sn}")
except Exception as e:
    print(f"Error: {e}")
